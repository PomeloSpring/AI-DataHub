"""会话工作区结构化文件预览与编辑 — Excel(xlsx) 多 sheet 读取与单元格级保存.

文件读写一律在会话工作区内(定位/防穿越/拒符号链接与 /api/chat/session-file 同一守卫);
访问控制为"属主自检"口径(与 workspace-assets 家族一致): 仅本人会话可读写。

- GET  /api/chat/session-file/xlsx       多 sheet 清单 + 指定 sheet 的表格数据(JSON)
- POST /api/chat/session-file/xlsx/save  单元格级修改保存(**生成新版本文件,原文件保留**)

保存不做原地覆盖的原因: openpyxl 重写可能丢失原文件中的图表/图片等高级特性,
原稿保留即天然回退路径(显式可诊断,不做静默覆盖)。
"""
import logging
import os
import re
from datetime import date, datetime, time
from decimal import Decimal

from fastapi import APIRouter, Depends, HTTPException, Query
from pydantic import BaseModel, Field

from services.shared.common.auth import get_file_user

logger = logging.getLogger(__name__)
router = APIRouter()

# 单 sheet 预览行数上限(超出截断并显式标注,前端网格虚拟滚动承载)
XLSX_MAX_ROWS = 5000
# 单次保存的修改单元格上限 / 单元格字符上限(Excel 单元格上限 32767)
MAX_CHANGES = 5000
MAX_CELL_CHARS = 32000
# Excel 物理边界
MAX_ROW = 1_048_576
MAX_COL = 16_384


def _resolve_session_ws(conversation_id: int, user_id: int):
    """会话 → 工作区目录;仅本人会话可访问。"""
    from services.datamind.execution import session_workspace as sw
    from services.shared.common.db.metadata_db import get_metadata_conn

    conn = get_metadata_conn()
    try:
        with conn.cursor() as cur:
            cur.execute(
                "SELECT session_key, workspace_id FROM adh_agent_sessions "
                "WHERE conversation_id=%s AND user_id=%s",
                (conversation_id, user_id))
            row = cur.fetchone()
    finally:
        conn.close()
    if not row or not row.get("session_key"):
        raise HTTPException(status_code=404, detail="会话不存在或无权访问")
    base = sw.workspace_base(create=False)
    root = sw.session_paths(base, row["session_key"], int(row["workspace_id"] or 0), create=False)
    return (root / "workspace").resolve()


def _resolve_target(ws_dir, path: str):
    """工作区内解析目标文件(防目录穿越/符号链接逃逸,与 session-file 伺服同一守卫)."""
    from services.datamind.multimodal.loader import normalize_rel_path, resolve_workspace_file

    if not normalize_rel_path(path):
        raise HTTPException(status_code=400, detail="缺少文件路径")
    try:
        target = resolve_workspace_file(ws_dir, path)
    except ValueError as e:
        raise HTTPException(status_code=403, detail=str(e) or "非法文件路径")
    if not target.is_file():
        raise HTTPException(status_code=404, detail="文件不存在")
    if target.suffix.lower() != ".xlsx":
        raise HTTPException(status_code=400, detail="仅支持 .xlsx 文件")
    return target


def _json_cell(v):
    """单元格值 → JSON 友好形态(公式经 data_only 取缓存值)."""
    if isinstance(v, (datetime, date, time)):
        return v.isoformat()
    if isinstance(v, Decimal):
        return float(v)
    if isinstance(v, bytes):
        return None
    return v


def _columns_from_header(header) -> list[str]:
    """表头 → 唯一列名: 空列名占位 col_N,重名列追加 _2/_3。"""
    columns: list[str] = []
    seen: dict[str, int] = {}
    for i, raw in enumerate(header or []):
        base = str(raw).strip() if raw is not None else ""
        if not base:
            base = f"col_{i + 1}"
        n = seen.get(base, 0)
        seen[base] = n + 1
        columns.append(base if n == 0 else f"{base}_{n + 1}")
    return columns


def _coerce_value(raw: str, original):
    """编辑值按原单元格类型智能转换: 数字/布尔原格尽量保持类型,空串清空。"""
    s = "" if raw is None else str(raw)
    if s == "":
        return None
    if isinstance(original, bool):
        return s.strip().lower() in ("true", "1", "yes", "y", "是")
    if isinstance(original, (int, float)) and not isinstance(original, bool):
        try:
            return int(s) if re.fullmatch(r"[+-]?\d+", s.strip()) else float(s.strip())
        except ValueError:
            return s
    return s


def _next_version_path(target) -> "os.PathLike":
    """生成新版本文件名: {stem}_edited.xlsx,冲突则 _edited_1.xlsx…(编辑对象本身是
    _edited 版本时,基名回溯到最初 stem,避免 _edited_edited 嵌套)."""
    stem = re.sub(r"_edited(_\d+)?$", "", target.stem)
    for n in range(0, 1000):
        name = f"{stem}_edited.xlsx" if n == 0 else f"{stem}_edited_{n}.xlsx"
        out = target.with_name(name)
        if not out.exists() and not out.is_symlink():
            return out
    raise HTTPException(status_code=409, detail="同名版本过多,请清理后重试")


class CellChange(BaseModel):
    row: int = Field(..., description="Excel 1-based 行号")
    col: int = Field(..., description="Excel 1-based 列号")
    value: str = Field("", description="单元格新值(按原格类型转换)")


class SaveXlsxRequest(BaseModel):
    conversation_id: int
    path: str
    sheet: str = ""
    changes: list[CellChange] = []


class SaveTableRequest(BaseModel):
    conversation_id: int
    filename: str = "查询结果.xlsx"
    columns: list[str] = []
    rows: list[dict] = []


MAX_TABLE_ROWS = 10000
MAX_TABLE_COLS = 200


@router.get("/session-file/xlsx")
def read_xlsx(
    conversation_id: int = Query(..., description="会话 ID"),
    path: str = Query(..., description="工作区相对路径"),
    sheet: str = Query("", description="工作表名,缺省第一张"),
    user: dict = Depends(get_file_user),
):
    """读取 xlsx: 工作表清单 + 指定 sheet 的表格数据(首行作表头)。"""
    ws_dir = _resolve_session_ws(conversation_id, int(user.get("user_id") or 0))
    target = _resolve_target(ws_dir, path)

    try:
        import openpyxl
        wb = openpyxl.load_workbook(target, read_only=True, data_only=True)
    except Exception as e:
        logger.exception("xlsx 解析失败: %s", target.name)
        raise HTTPException(status_code=422, detail=f"Excel 文件解析失败: {type(e).__name__}") from None

    try:
        names = list(wb.sheetnames)
        active = sheet if sheet in names else (names[0] if names else "")
        if not active:
            return {"sheets": [], "active": "", "columns": [], "rows": [], "row_count": 0, "truncated": False}
        ws = wb[active]
        rows_iter = ws.iter_rows(values_only=True)
        header = next(rows_iter, None)
        columns = _columns_from_header(header)
        rows = []
        truncated = False
        for values in rows_iter:
            if len(rows) >= XLSX_MAX_ROWS:
                truncated = True
                break
            if values is None or all(v is None for v in values):
                continue
            rows.append({c: _json_cell(values[i]) if i < len(values) else None
                         for i, c in enumerate(columns)})
        return {
            "sheets": names,
            "active": active,
            "columns": columns,
            "rows": rows,
            "row_count": len(rows),
            "truncated": truncated,
        }
    finally:
        wb.close()


@router.post("/session-file/xlsx/save")
def save_xlsx(req: SaveXlsxRequest, user: dict = Depends(get_file_user)):
    """单元格级保存:**写入新版本文件(原文件保留)**,返回新文件引用。

    修改按原格类型转换(数字/布尔保持类型,空串清空);不支持的坐标/超限显式拒绝。
    """
    if not req.changes:
        raise HTTPException(status_code=400, detail="没有需要保存的修改")
    if len(req.changes) > MAX_CHANGES:
        raise HTTPException(status_code=400, detail=f"单次最多保存 {MAX_CHANGES} 处修改")
    for ch in req.changes:
        if not (1 <= ch.row <= MAX_ROW) or not (1 <= ch.col <= MAX_COL):
            raise HTTPException(status_code=400, detail=f"单元格坐标超出 Excel 范围: R{ch.row}C{ch.col}")
        if len(ch.value or "") > MAX_CELL_CHARS:
            raise HTTPException(status_code=400, detail="单元格内容过长(上限 32000 字符)")

    ws_dir = _resolve_session_ws(req.conversation_id, int(user.get("user_id") or 0))
    target = _resolve_target(ws_dir, req.path)

    try:
        import openpyxl
        wb = openpyxl.load_workbook(target)
    except Exception as e:
        logger.exception("xlsx 打开失败: %s", target.name)
        raise HTTPException(status_code=422, detail=f"Excel 文件打开失败: {type(e).__name__}") from None

    if not req.sheet or req.sheet not in wb.sheetnames:
        raise HTTPException(status_code=400, detail="工作表不存在")
    ws = wb[req.sheet]
    try:
        for ch in req.changes:
            cell = ws.cell(row=ch.row, column=ch.col)
            cell.value = _coerce_value(ch.value, cell.value)
        out = _next_version_path(target)
        wb.save(out)
    except HTTPException:
        raise
    except Exception as e:
        logger.exception("xlsx 保存失败: %s", target.name)
        raise HTTPException(status_code=500, detail=f"Excel 保存失败: {type(e).__name__}") from None
    finally:
        wb.close()

    return {
        "filename": out.name,
        "path": str(out.relative_to(ws_dir)),
        "size": out.stat().st_size,
        "sheet": req.sheet,
    }


def _plain_cell(v):
    """结果表单元格 → openpyxl 可写形态(数字/布尔/文本保型,空→None,复杂值转字符串)."""
    if v is None or isinstance(v, (int, float, bool, str)):
        return v
    return str(v)


@router.post("/session-file/table")
def save_table(req: SaveTableRequest, user: dict = Depends(get_file_user)):
    """把结果表保存为会话工作区 xlsx 文件(文件管理工具口径: 落盘 uploads/,随会话清理)."""
    if not req.columns:
        raise HTTPException(status_code=400, detail="缺少表头列")
    if len(req.columns) > MAX_TABLE_COLS:
        raise HTTPException(status_code=400, detail=f"列数超过上限 {MAX_TABLE_COLS}")
    if len(req.rows) > MAX_TABLE_ROWS:
        raise HTTPException(status_code=400, detail=f"行数超过上限 {MAX_TABLE_ROWS}")

    ws_dir = _resolve_session_ws(req.conversation_id, int(user.get("user_id") or 0))

    import io
    import openpyxl
    from services.datamind.multimodal.loader import write_upload_file

    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.title = "Sheet1"
    sheet.append([str(c) for c in req.columns])
    for r in req.rows:
        sheet.append([_plain_cell(r.get(c)) for c in req.columns])
    buf = io.BytesIO()
    wb.save(buf)

    filename = req.filename if req.filename.lower().endswith(".xlsx") else f"{req.filename}.xlsx"
    placed = write_upload_file(ws_dir, filename, buf.getvalue())
    return {
        "filename": placed["filename"],
        "path": placed["path"],
        "size": placed["size"],
        "row_count": len(req.rows),
    }
