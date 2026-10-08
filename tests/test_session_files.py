"""会话工作区 xlsx 预览/编辑 —— 离线单测(假 DB + 临时工作区,不触碰真实库)。

覆盖:
  S1  读取: 多 sheet 清单/首行表头/数值与日期 JSON 友好化/空行跳过;
  S2  读取守卫: 目录穿越 403、非 xlsx 400、无会话 404;
  S3  保存: 单元格级修改按原格类型转换,生成新版本文件且原文件保留;
  S4  保存守卫: 空修改 400、坐标越界 400、超量 400;
  S5  版本命名: 对 _edited 版本继续保存不产生 _edited_edited 嵌套;
  S6  表头处理: 空列名占位 col_N、重名列去重。
"""
import datetime
import importlib
import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

import pytest
from fastapi import HTTPException

sf = importlib.import_module("backend.modules.mind.api.session_files")
sw = importlib.import_module("backend.modules.mind.execution.session_workspace")
metadata_db = importlib.import_module("backend.common.db.metadata_db")

USER = {"user_id": 7, "username": "tester", "role": "analyst"}
SESSION_KEY = "a" * 32


class _FakeCursor:
    def __init__(self, row):
        self.row = row

    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False

    def execute(self, sql, params=None):
        self._sql = sql

    def fetchone(self):
        return self.row


class _FakeConn:
    def __init__(self, row):
        self.row = row

    def cursor(self):
        return _FakeCursor(self.row)

    def close(self):
        pass


@pytest.fixture
def ws(tmp_path, monkeypatch):
    """临时会话工作区 + 假 DB(仅本人会话)。"""
    monkeypatch.setattr("backend.common.config.ADH_WORKSPACES_DIR", str(tmp_path))
    monkeypatch.setattr(metadata_db, "get_metadata_conn",
                        lambda: _FakeConn({"session_key": SESSION_KEY, "workspace_id": 3}))
    root = sw.session_paths(tmp_path, SESSION_KEY, 3, create=True)
    return root / "workspace"


def _write_xlsx(ws_dir, name="a.xlsx", second=False):
    import openpyxl
    (ws_dir / "uploads").mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.title = "Sheet1"
    sheet.append(["名称", "数量", "日期"])
    sheet.append(["alpha", 1, datetime.date(2026, 10, 3)])
    sheet.append(["beta", 2.5, None])
    if second:
        s2 = wb.create_sheet("Second")
        s2.append(["x"])
        s2.append([9])
    out = ws_dir / "uploads" / name
    wb.save(out)
    return out


# ── S1 读取 ────────────────────────────────────────────────────────

def test_read_basic(ws):
    _write_xlsx(ws, second=True)
    r = sf.read_xlsx(conversation_id=1, path="uploads/a.xlsx", sheet="", user=USER)
    assert r["sheets"] == ["Sheet1", "Second"]
    assert r["active"] == "Sheet1"
    assert r["columns"] == ["名称", "数量", "日期"]
    assert r["row_count"] == 2 and r["truncated"] is False
    assert r["rows"][0]["名称"] == "alpha" and r["rows"][0]["数量"] == 1
    # 日期 → ISO 字符串(openpyxl 将日期格读回 datetime,故可能带 T00:00:00),空单元格 → None
    assert str(r["rows"][0]["日期"]).startswith("2026-10-03")
    assert r["rows"][1]["日期"] is None
    assert r["rows"][1]["数量"] == 2.5


def test_read_sheet_switch(ws):
    _write_xlsx(ws, second=True)
    r = sf.read_xlsx(conversation_id=1, path="uploads/a.xlsx", sheet="Second", user=USER)
    assert r["active"] == "Second"
    assert r["columns"] == ["x"] and r["rows"] == [{"x": 9}]


def test_read_skips_blank_rows(ws):
    import openpyxl
    (ws_dir := ws / "uploads").mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    wb.active.append(["k", "v"])
    wb.active.append(["a", 1])
    wb.active.append([None, None])
    wb.active.append(["b", 2])
    wb.save(ws_dir / "blank.xlsx")
    r = sf.read_xlsx(conversation_id=1, path="uploads/blank.xlsx", sheet="", user=USER)
    assert r["row_count"] == 2


# ── S2 读取守卫 ────────────────────────────────────────────────────

def test_read_rejects_traversal(ws):
    _write_xlsx(ws)
    with pytest.raises(HTTPException) as e:
        sf.read_xlsx(conversation_id=1, path="../../etc/passwd", sheet="", user=USER)
    assert e.value.status_code == 403


def test_read_rejects_non_xlsx(ws):
    (ws / "uploads").mkdir(parents=True, exist_ok=True)
    (ws / "uploads" / "a.txt").write_text("hi", encoding="utf-8")
    with pytest.raises(HTTPException) as e:
        sf.read_xlsx(conversation_id=1, path="uploads/a.txt", sheet="", user=USER)
    assert e.value.status_code == 400


def test_read_missing_session(ws, monkeypatch):
    monkeypatch.setattr(metadata_db, "get_metadata_conn", lambda: _FakeConn(None))
    with pytest.raises(HTTPException) as e:
        sf.read_xlsx(conversation_id=1, path="uploads/a.xlsx", sheet="", user=USER)
    assert e.value.status_code == 404


# ── S3 保存 ────────────────────────────────────────────────────────

def test_save_writes_new_version_and_preserves_original(ws):
    import openpyxl
    _write_xlsx(ws)
    req = sf.SaveXlsxRequest(
        conversation_id=1, path="uploads/a.xlsx", sheet="Sheet1",
        changes=[sf.CellChange(row=2, col=2, value="42"),   # 数字格保持数字
                 sf.CellChange(row=3, col=1, value=""),     # 空串清空
                 sf.CellChange(row=3, col=3, value="2026-10-04")],  # 非数字格写字符串
    )
    r = sf.save_xlsx(req, user=USER)
    assert r["filename"] == "a_edited.xlsx"
    assert r["path"] == "uploads/a_edited.xlsx"

    new = openpyxl.load_workbook(ws / "uploads" / "a_edited.xlsx")
    s = new["Sheet1"]
    assert s.cell(row=2, column=2).value == 42
    assert s.cell(row=3, column=1).value is None
    assert s.cell(row=3, column=3).value == "2026-10-04"

    # 原文件保持不变
    old = openpyxl.load_workbook(ws / "uploads" / "a.xlsx")
    assert old["Sheet1"].cell(row=2, column=2).value == 1


# ── S4 保存守卫 ────────────────────────────────────────────────────

def test_save_rejects_empty_changes(ws):
    _write_xlsx(ws)
    req = sf.SaveXlsxRequest(conversation_id=1, path="uploads/a.xlsx", sheet="Sheet1", changes=[])
    with pytest.raises(HTTPException) as e:
        sf.save_xlsx(req, user=USER)
    assert e.value.status_code == 400


def test_save_rejects_out_of_range(ws):
    _write_xlsx(ws)
    req = sf.SaveXlsxRequest(
        conversation_id=1, path="uploads/a.xlsx", sheet="Sheet1",
        changes=[sf.CellChange(row=0, col=1, value="x")])
    with pytest.raises(HTTPException) as e:
        sf.save_xlsx(req, user=USER)
    assert e.value.status_code == 400


def test_save_rejects_too_many_changes(ws):
    _write_xlsx(ws)
    req = sf.SaveXlsxRequest(
        conversation_id=1, path="uploads/a.xlsx", sheet="Sheet1",
        changes=[sf.CellChange(row=2, col=1, value="x")] * (sf.MAX_CHANGES + 1))
    with pytest.raises(HTTPException) as e:
        sf.save_xlsx(req, user=USER)
    assert e.value.status_code == 400


def test_save_rejects_unknown_sheet(ws):
    _write_xlsx(ws)
    req = sf.SaveXlsxRequest(
        conversation_id=1, path="uploads/a.xlsx", sheet="Nope",
        changes=[sf.CellChange(row=2, col=1, value="x")])
    with pytest.raises(HTTPException) as e:
        sf.save_xlsx(req, user=USER)
    assert e.value.status_code == 400


# ── S5 版本命名 ────────────────────────────────────────────────────

def test_save_again_no_nested_edited_suffix(ws):
    _write_xlsx(ws)
    req = sf.SaveXlsxRequest(
        conversation_id=1, path="uploads/a.xlsx", sheet="Sheet1",
        changes=[sf.CellChange(row=2, col=1, value="v1")])
    r1 = sf.save_xlsx(req, user=USER)
    assert r1["filename"] == "a_edited.xlsx"

    req2 = sf.SaveXlsxRequest(
        conversation_id=1, path=r1["path"], sheet="Sheet1",
        changes=[sf.CellChange(row=2, col=1, value="v2")])
    r2 = sf.save_xlsx(req2, user=USER)
    # 基名回溯,不产生 a_edited_edited.xlsx
    assert r2["filename"] == "a_edited_1.xlsx"
    assert (ws / "uploads" / "a_edited.xlsx").exists()


# ── S6 表头处理 ────────────────────────────────────────────────────

def test_header_placeholder_and_dedupe(ws):
    import openpyxl
    (ws / "uploads").mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    wb.active.append(["a", None, "a"])
    wb.active.append([1, 2, 3])
    wb.save(ws / "uploads" / "h.xlsx")
    r = sf.read_xlsx(conversation_id=1, path="uploads/h.xlsx", sheet="", user=USER)
    assert r["columns"] == ["a", "col_2", "a_2"]
    assert r["rows"] == [{"a": 1, "col_2": 2, "a_2": 3}]


# ── S7 结果表落盘工作区 ────────────────────────────────────────

def test_save_table_creates_xlsx(ws):
    import openpyxl
    req = sf.SaveTableRequest(
        conversation_id=1, filename="查询结果.xlsx",
        columns=["名称", "数值"],
        rows=[{"名称": "a", "数值": 1}, {"名称": "b", "数值": 2.5}],
    )
    r = sf.save_table(req, user=USER)
    assert r["path"] == "uploads/查询结果.xlsx"
    assert r["row_count"] == 2

    wb = openpyxl.load_workbook(ws / r["path"])
    s = wb.active
    assert [s.cell(row=1, column=1).value, s.cell(row=1, column=2).value] == ["名称", "数值"]
    assert s.cell(row=2, column=2).value == 1  # 数字保型
    assert s.cell(row=3, column=2).value == 2.5
    # 生成的文件可被读取端完整解析(落盘↔预览闭环)
    back = sf.read_xlsx(conversation_id=1, path=r["path"], sheet="", user=USER)
    assert back["rows"] == [{"名称": "a", "数值": 1}, {"名称": "b", "数值": 2.5}]


def test_save_table_dedupe_and_suffix(ws):
    req = sf.SaveTableRequest(
        conversation_id=1, filename="结果", columns=["a"], rows=[{"a": 1}])
    r1 = sf.save_table(req, user=USER)
    assert r1["filename"] == "结果.xlsx"
    r2 = sf.save_table(req, user=USER)
    assert r2["filename"] == "结果_1.xlsx"


def test_save_table_rejects_bad_input(ws):
    with pytest.raises(HTTPException) as e:
        sf.save_table(sf.SaveTableRequest(conversation_id=1, columns=[], rows=[]), user=USER)
    assert e.value.status_code == 400

    big = sf.SaveTableRequest(
        conversation_id=1, columns=["a"],
        rows=[{"a": 1}] * (sf.MAX_TABLE_ROWS + 1))
    with pytest.raises(HTTPException) as e:
        sf.save_table(big, user=USER)
    assert e.value.status_code == 400
